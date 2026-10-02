"""수주 파서의 헬퍼 함수 — 날짜/숫자 변환"""
import sys, os
from datetime import date, datetime
sys.path.insert(0, os.path.dirname(os.path.dirname(__file__)))

from utils.so_parser import _to_num, _to_int, _to_date, _pick_line_nums


def test_to_num_basic():
    assert _to_num("1,234") == 1234.0
    assert _to_num("1234") == 1234.0
    assert _to_num(1234) == 1234.0
    assert _to_num(None) is None
    assert _to_num("") is None
    assert _to_num("abc") is None


def test_to_int_basic():
    assert _to_int("100") == 100
    assert _to_int("1,234") == 1234
    assert _to_int(None) is None


def test_pick_line_nums_normal_line_not_merged():
    """정상 라인은 병합 금지 — MJT-PO26-우성-708 사고 (100 이
    6,000·600,000 과 이어붙어 1조가 되던 버그) 회귀 방지"""
    assert _pick_line_nums(["100", "6,000", "600,000"]) == \
        [100, 6000, 600000]


def test_pick_line_nums_split_amount_merged():
    """PDF 가 쪼갠 금액("7 8,000,000")은 검산으로 복원"""
    assert _pick_line_nums(["13,000", "6,000", "7", "8,000,000"]) == \
        [13000, 6000, 78000000]


def test_pick_line_nums_two_tokens_kept():
    assert _pick_line_nums(["50", "1,200"]) == [50, 1200]


def test_pick_line_nums_double_split_merged():
    """수량·금액이 한 행에서 동시에 쪼개진 경우 — MJT-PO26-우성-725
    ("3 ,600 6,000 2 1,600,000" = 3,600 × 6,000 = 21,600,000).
    단일 병합만 시도하던 버그는 3 / 600 / 6,000 으로 읽었다"""
    assert _pick_line_nums(["3", ",600", "6,000", "2", "1,600,000"]) == \
        [3600, 6000, 21600000]


def test_to_date_formats():
    assert _to_date("2026-05-08") == date(2026, 5, 8)
    assert _to_date("2026/05/08") == date(2026, 5, 8)
    assert _to_date("20260508") == date(2026, 5, 8)
    assert _to_date(datetime(2026, 5, 8, 13, 30)) == date(2026, 5, 8)
    assert _to_date(None) is None
    assert _to_date("") is None
    assert _to_date("not a date") is None


# ─── DIC 납품예정서등록 엑셀 (2026-10-02) ───
def _dic_workbook():
    import io as _io
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "납품예정서등록(DIC)"
    ws["A2"] = "납품예정서등록-DIC"
    ws["A5"] = "납품예정서등록(DIC)"
    ws.append([])
    ws["A10"], ws["B10"], ws["C10"] = "업체", "36930", "우성정밀"
    hdr = [None, "수주일자", "품목", "품목명", "규격", "재질", "구분", "납품공장",
           "납품공장명", "단위", "수주번호", "행번", "수주량", "납기일", "납품량",
           "미납품량", "납품창고명"]
    for j, h in enumerate(hdr, 1):
        ws.cell(17, j, h)
    data = [
        [1, "2026-09-15", "6A433004#1", "COLLOR,SPLINE-선삭", "Ø99*37", "SCM415",
         "정상", "P10", "중장비공장", "EA ", "PO2609000000712", "3", 400,
         "2026-10-31", 382, 18, "통합물류창고(중장비)"],
        [2, "2026-10-01", "624315-45001#1", "SHAFT-IDLER-선삭", "", "S45C",
         "정상", "P20", "차량2공장", "EA ", "PO2610000000060", "1", 800,
         "2026-10-09", 0, 800, "통합물류창고(차량2)"],
        [3, "2026-10-01", "X-1", "신규 공장 품목", "", "", "정상", "P99",
         "신공장", "EA", "PO2610000000999", "1", 10, "2026-10-09", 0, 10, ""],
    ]
    for i, row in enumerate(data):
        for j, v in enumerate(row, 1):
            ws.cell(18 + i, j, v)
    # 아래쪽 다른 표 머리글 — 데이터로 읽으면 안 됨
    for j, h in enumerate([None, "Title", "납품예정일자", "납품예정수량"], 1):
        ws.cell(23, j, h)
    buf = _io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_dic_excel_detected_and_parsed():
    from datetime import date
    from utils.so_parser import (
        detect_so_format, group_by_so_number, parse_so_auto)
    raw = _dic_workbook()
    assert detect_so_format(raw, "납품예정서등록-DIC.xls") == "DIC"
    fmt, items = parse_so_auto(raw, "납품예정서등록-DIC.xls")
    assert fmt == "DIC" and len(items) == 3
    a, b, c = items
    # 납품공장 → 거래처 마스터 정식명 (중장비 = 두서, 차량2 = 두동)
    assert a["customer"] == "(주)디아이씨두서공장"
    assert b["customer"] == "(주)디아이씨두동공장"
    assert "신공장" in c["customer"]          # 모르는 공장은 그대로 드러낸다
    assert a["so_number"] == "PO2609000000712" and a["line_no"] == 3
    assert a["customer_part_no"] == "6A433004#1"
    # 미진 양식과 같은 형식 — 수량 = 수주량(계약 수량), 납품량 = 기납품 → 미납 18
    assert a["qty"] == 400 and a["received_qty"] == 382
    assert "원 수주량" not in (a["remark"] or "")
    assert a["so_date"] == date(2026, 9, 15) and a["due_date"] == date(2026, 10, 31)
    assert a["unit"] == "EA" and a["unit_price"] is None  # 단가 열 없음
    assert "Ø99*37" in a["remark"] and "SCM415" in a["remark"]
    # 한 파일에 거래처가 둘 — (거래처, 수주번호) 로 묶인다
    heads = {(g["header"]["customer"], g["header"]["so_number"])
             for g in group_by_so_number(items)}
    assert ("(주)디아이씨두서공장", "PO2609000000712") in heads
    assert ("(주)디아이씨두동공장", "PO2610000000060") in heads


def test_hdx_excel_reads_received_qty_as_pre_delivered():
    """HDX 양식도 미진·DIC 와 같은 형식 — 수량 = 수주 수량, 입고수량 = 기납품."""
    import io as _io
    from datetime import datetime
    import openpyxl
    from utils.so_parser import parse_so_auto
    wb = openpyxl.Workbook()
    ws = wb.active
    hdr = ["No", "MRP", "입고상태", "수주번호", "항번", "수주일자", "업체자재코드",
           "납기요청일", "협력업체", "자재코드", "자재명", "수량", "단위", "단가",
           "금액", "납품처 장소", "입고수량", "특이사항"]
    ws.append(hdr)
    ws.append([1, "비대상", "미입고", "G26A920021", "00001",
               datetime(2026, 10, 1), "", datetime(2026, 10, 16), "우성정밀",
               "HA30-80730", "허브 로크 너트", 500, "EA", 7542, 3771000,
               "Ulsan", "", ""])
    ws.append([2, "비대상", "부분입고", "G26A920021", "00002",
               datetime(2026, 10, 1), "", datetime(2026, 10, 30), "우성정밀",
               "HA80-20210", "주차 디스크", 100, "EA", 51935, 5193500,
               "Ulsan", 40, ""])
    buf = _io.BytesIO()
    wb.save(buf)
    fmt, items = parse_so_auto(buf.getvalue(), "2026.10.02_15.43.35.xlsx")
    assert fmt == "HDX" and len(items) == 2
    assert items[0]["customer"] == "HDX" and items[0]["line_no"] == 1
    assert items[0]["qty"] == 500 and items[0]["received_qty"] == 0
    assert items[1]["qty"] == 100 and items[1]["received_qty"] == 40
    assert items[1]["unit_price"] == 51935
