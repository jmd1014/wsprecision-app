"""핫픽스 3(2026-10-30) — 미진 엑셀 업로드 거래처 판정.

`parse_mijin_excel` 이 모든 라인을 "미진정밀" 로 적어 엠제이티 전용 품번(MJT_ONLY_PNS) 라인이 미진정밀
수주로 들어갔다. ERP 파서(group_by_order)와 같은 판정으로 `(주)엠제이티` + 수주번호 `-MJT` 로 나눈다.
- 파서: 같은 외주발주번호에 미진·엠제이티 품목이 섞이면 수주 2개(line_no 는 수주별 1부터), 금액 분리
- 대조군: 순수 미진 파일은 결과 불변
- `_strip_prefix` → utils.so_parser.strip_prefix 이동: 예전 함수와 같은 결과(S/4S/SFB·S630 예외 포함)
- 화면: 중복 검증이 ((주)엠제이티, X-MJT) 키로 기존 수주를 찾고, 개정 발주 대조도 거래처별
"""
import importlib.util
import io
import os
import sys
from urllib.parse import unquote

import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
from tests.conftest import goto  # noqa: E402
from utils.so_parser import group_by_so_number, parse_mijin_excel, strip_prefix  # noqa: E402

APP_FILE = os.path.join(ROOT, "streamlit_app.py")
app_test = pytest.mark.skipif(importlib.util.find_spec("streamlit") is None,
                              reason="streamlit 미설치 — AppTest 불가")

PO = "OP2610300001"
MIJIN, MJT = "미진정밀", "(주)엠제이티"
HEAD = ["선택", "외주발주번호", "외주발주일", "납기일", "제품번호", "제품명",
        "발주수량", "납품수량", "단가", "금액", "부가세", "금액계"]
# (제품번호, 제품명, 수량, 단가)
MIXED = [("4S4PDVN-03;OUP", "PDVN 3", 100, 1200),
         ("4HDV-VM-04;OUP", "HDV 4", 50, 3000),
         ("4S8PDVN-02;OUP", "PDVN 8", 20, 5000)]
PURE = [("4HDV-VM-04;OUP", "HDV 4", 50, 3000),
        ("S12HDV-VM-01;OUP", "HDV 1", 10, 700)]


def _workbook(rows, po=PO):
    """행 1 타이틀 · 행 2 헤더 · 행 3 TOTAL · 행 4~ 데이터(미진 외주발주품목조회 양식)."""
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.cell(1, 1, "외주발주품목조회")
    for j, h in enumerate(HEAD, 1):
        ws.cell(2, j, h)
    ws.cell(3, 2, "TOTAL")
    for i, (pn, name, q, p) in enumerate(rows):
        r = 4 + i
        for j, v in enumerate(["N", po, "2026-10-28", None, pn, name, q, 0, p, q * p,
                               q * p // 10, q * p * 11 // 10], 1):
            ws.cell(r, j, v)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── 파서 ─────────────────────────────────────────────────────────────

def test_mixed_po_splits_into_mijin_and_mjt_orders():
    items = parse_mijin_excel(_workbook(MIXED), "외주발주품목조회.xlsx")
    assert [(it["customer"], it["so_number"], it["line_no"]) for it in items] == [
        (MJT, f"{PO}-MJT", 1), (MIJIN, PO, 1), (MJT, f"{PO}-MJT", 2)]
    groups = group_by_so_number(items)
    by = {(g["header"]["customer"], g["header"]["so_number"]): g for g in groups}
    assert set(by) == {(MIJIN, PO), (MJT, f"{PO}-MJT")}
    mj, mt = by[(MIJIN, PO)], by[(MJT, f"{PO}-MJT")]
    assert [it["customer_part_no"] for it in mj["items"]] == ["4HDV-VM-04;OUP"]
    assert [(it["customer_part_no"], it["line_no"]) for it in mt["items"]] == [
        ("4S4PDVN-03;OUP", 1), ("4S8PDVN-02;OUP", 2)]
    assert mj["header"]["total_amount"] == 50 * 3000
    assert mt["header"]["total_amount"] == 100 * 1200 + 20 * 5000
    # 원본 품번·힌트는 그대로(매칭은 alias 로)
    assert mt["items"][0]["canonical_pn_hint"] == "4S4PDVN-03;OUP"


def test_mjt_detection_ignores_case_and_other_prefixes():
    rows = [("4s4pdvn-02;oup", "a", 1, 1), ("SFB-16;PM", "b", 1, 1), ("S12HFDVN-VM-03", "c", 1, 1),
            ("8HFDV-VM-05;OUP", "d", 1, 1)]     # 8HFDV-VM-05 는 목록에 없음 → 미진(지시문 '하지 않는 것')
    items = parse_mijin_excel(_workbook(rows), "x.xlsx")
    assert [it["customer"] for it in items] == [MJT, MJT, MJT, MIJIN]


def test_pure_mijin_file_unchanged():
    """대조군 — 엠제이티 품번이 없으면 예전과 같다(묶음 전 커밋에서도 통과)."""
    items = parse_mijin_excel(_workbook(PURE), "x.xlsx")
    assert [(it["customer"], it["so_number"], it["line_no"], it["customer_part_no"], it["qty"],
             it["amount"]) for it in items] == [
        (MIJIN, PO, 1, "4HDV-VM-04;OUP", 50.0, 150000.0),
        (MIJIN, PO, 2, "S12HDV-VM-01;OUP", 10.0, 7000.0)]
    groups = group_by_so_number(items)
    assert len(groups) == 1 and groups[0]["header"]["customer"] == MIJIN
    assert groups[0]["header"]["so_number"] == PO and groups[0]["header"]["total_amount"] == 157000


# ── strip_prefix 이동 ─────────────────────────────────────────────────

def _old_strip_prefix(s):
    """이동 전 streamlit_app 지역 함수(_strip_prefix·_so_strip_prefix) 그대로."""
    if not s: return ""
    p = str(s).upper().strip()
    if p.startswith('4S') and len(p) > 2 and (p[2].isalnum()):
        return p[2:]
    if p.startswith('S') and len(p) > 1 and (p[1].isalnum()):
        excluded = ('SP-', 'SDF', 'SUS', 'SODV', 'SFB', 'SCM', 'SKD', 'SKH',
                    'S45', 'S20', 'S30', 'S304', 'S316', 'S630')
        if not any(p.startswith(e) for e in excluded):
            return p[1:]
    return p


@pytest.mark.parametrize("s,want", [
    ("4S4PDVN-03", "4PDVN-03"), ("4s8pdvn-02", "8PDVN-02"), ("S12HFDVN-VM-03", "12HFDVN-VM-03"),
    ("SFB-16", "SFB-16"), ("4SSFB-24", "SFB-24"), ("S630-PIN", "S630-PIN"), ("SUS304-A", "SUS304-A"),
    ("SP-100", "SP-100"), ("SODV-1", "SODV-1"), ("S45C-R", "S45C-R"), ("4S", "4S"), ("4S-1", "4S-1"),
    ("S", "S"), (" 4S4HDV-VM-04 ", "4HDV-VM-04"), ("", ""), (None, ""),
])
def test_strip_prefix_same_as_before(s, want):
    assert strip_prefix(s) == want == _old_strip_prefix(s)


def test_app_uses_utils_strip_prefix_only():
    src = open(APP_FILE, encoding="utf-8").read()
    assert "def _strip_prefix(" not in src and "def _so_strip_prefix(" not in src
    assert "strip_prefix as _so_strip_prefix" in src
    assert 'customer_name = items[0]["customer"]' not in src       # 단일 거래처 가정 제거


# ── 화면(AppTest, 목 DB) ────────────────────────────────────────────

PRODUCTS = [
    {"product_id": "P403", "pn": "4PDVN-03", "alias_list": "4S4PDVN-03", "archived_at": None,
     "archive_reason": None},
    {"product_id": "P802", "pn": "8PDVN-02", "alias_list": "4S8PDVN-02", "archived_at": None,
     "archive_reason": None},
    {"product_id": "P404", "pn": "4HDV-VM-04", "alias_list": None, "archived_at": None,
     "archive_reason": None},
    {"product_id": "P402", "pn": "4PDVN-02", "alias_list": "OLD-PDVN02", "archived_at": None,
     "archive_reason": None},
]
INSERTED, FETCHES = [], []


def _mk_fetch(existing_mjt):
    """existing_mjt: ((주)엠제이티, PO-MJT) 수주가 이미 있다(라인 2개, 파일과 같은 수량·단가)."""
    def _fetch(table, select="*", filter_query="", limit=1000, **kw):
        fq = unquote(filter_query or "")
        FETCHES.append((table, fq))
        if table == "vendors":
            if f'name=eq."{MJT}"' in fq:
                return [{"vendor_id": 22, "name": MJT}]
            if f'name=eq."{MIJIN}"' in fq:
                return [{"vendor_id": 73, "name": MIJIN}]
            return []                                    # short_name 은 둘 다 NULL
        if table == "sales_orders":
            if existing_mjt and f'customer=eq."{MJT}"' in fq and f"{PO}-MJT" in fq \
                    and "so_number=in." in fq:
                return [{"so_id": 600, "so_number": f"{PO}-MJT"}]
            return []
        if table == "sales_order_items":
            if existing_mjt and "so_id=in.(600)" in fq:
                return [{"so_id": 600, "soi_id": 1, "line_no": 1, "customer_part_no": "4S4PDVN-03;OUP",
                         "qty": 100, "unit_price": 1200},
                        {"so_id": 600, "soi_id": 2, "line_no": 2, "customer_part_no": "4S8PDVN-02;OUP",
                         "qty": 20, "unit_price": 5000}]
            return []
        if table == "products":
            return [] if "product_id=in." in fq else [dict(p) for p in PRODUCTS]
        return []
    return _fetch


def _fetch_one(table, filter_query="", select="*"):
    fq = unquote(filter_query or "")
    if table == "sales_orders":
        if f"{PO}-MJT" in fq:
            return {"so_id": 702}
        if PO in fq:
            return {"so_id": 701}
    return None


@pytest.fixture()
def up_db(monkeypatch):
    import db
    import utils.slack_notify as _sn
    INSERTED.clear()
    FETCHES.clear()
    monkeypatch.setattr(_sn, "notify", lambda text, sync=False: True)
    monkeypatch.setattr(db, "fetch_one", _fetch_one)
    monkeypatch.setattr(db, "insert", lambda t, r: (INSERTED.append((t, r)), len(r))[1])
    monkeypatch.setattr(db, "update", lambda *a, **k: True)
    monkeypatch.setattr(db, "health_check", lambda: {"status": "OK", "counts": {}})
    monkeypatch.setattr(db, "debug_check", lambda: {"status": "mock"})
    monkeypatch.setattr(db, "count_rows", lambda *a, **k: 0)
    monkeypatch.setattr(db, "rpc", lambda *a, **k: None)

    def _use(existing_mjt):
        monkeypatch.setattr(db, "fetch", _mk_fetch(existing_mjt))
        return db
    return _use


def _upload(rows):
    import streamlit as st
    from streamlit.testing.v1 import AppTest
    for _c in (st.cache_data, st.cache_resource):
        try:
            _c.clear()
        except Exception:
            pass
    at = AppTest.from_file(APP_FILE, default_timeout=60)
    at.secrets["supabase"] = {"url": "https://mock.local", "anon_key": "a", "service_role_key": "s"}
    at.secrets["auth"] = {"disabled": True}
    at.run()
    goto(at, "so", "input")
    at.get("file_uploader")[0].upload("외주발주품목조회.xlsx", _workbook(rows)).run()
    assert not at.exception, [str(e.value) for e in at.exception]
    return at


def _save(at):
    next(b for b in at.button if b.key == "so_up_save").click().run()
    assert not at.exception, [str(e.value) for e in at.exception]
    heads = [r for t, rows in INSERTED if t == "sales_orders" for r in rows]
    lines = [r for t, rows in INSERTED if t == "sales_order_items" for r in rows]
    return heads, lines


@app_test
def test_upload_new_mixed_file_saves_two_orders(up_db):
    up_db(existing_mjt=False)
    at = _upload(MIXED)
    from tests.conftest import msg_texts
    assert "엠제이티 분리 1건" in " ".join(msg_texts(at))
    assert not any("거래처 판정이 마스터 품번과 다른" in str(w.value) for w in at.warning)
    heads, lines = _save(at)
    hd = {(h["customer"], h["so_number"]): h for h in heads}
    assert set(hd) == {(MIJIN, PO), (MJT, f"{PO}-MJT")}
    assert hd[(MJT, f"{PO}-MJT")]["vendor_id"] == 22 and hd[(MIJIN, PO)]["vendor_id"] == 73
    # alias(4S4PDVN-03·4S8PDVN-02)로 마스터 연결 — 거래처가 바뀌어도 매칭 유지
    mt = sorted((x["line_no"], x["product_id"]) for x in lines if x["so_id"] == 702)
    assert mt == [(1, "P403"), (2, "P802")]
    assert [(x["line_no"], x["product_id"]) for x in lines if x["so_id"] == 701] == [(1, "P404")]


@app_test
@pytest.mark.parametrize("rows", [MIXED, [MIXED[1], MIXED[0], MIXED[2]]], ids=["mjt_first", "mijin_first"])
def test_reupload_finds_existing_mjt_order_by_customer_key(up_db, rows):
    """중복 검증이 ((주)엠제이티, X-MJT) 키로 기존 수주를 찾아 제외 — 개정 발주(단가·수량 다른 라인)로도
    잡지 않는다(거래처별 대조). 미진 쪽 새 수주 X 만 저장. mijin_first: 예전 개정 대조는 첫 줄 거래처
    (customer_name = items[0]) 하나로만 찾아, 미진 줄이 앞이면 엠제이티 라인을 '다른 라인'으로 잡았다."""
    up_db(existing_mjt=True)
    at = _upload(rows)
    assert any(f'customer=eq."{MJT}"' in fq and f"so_number=in.({PO}-MJT)" in fq
               for t, fq in FETCHES if t == "sales_orders")
    assert any("이미 등록된 수주 **1건**" in str(w.value) and f"{PO}-MJT" in str(w.value)
               for w in at.warning), [w.value for w in at.warning]
    from tests.conftest import msg_texts
    assert not any("단가·수량이 다른 라인" in t for t in msg_texts(at))
    assert not [c for c in at.checkbox if c.key == "so_rev_ok"]
    heads, lines = _save(at)
    assert [(h["customer"], h["so_number"]) for h in heads] == [(MIJIN, PO)]
    assert [(x["so_id"], x["product_id"]) for x in lines] == [(701, "P404")]


@app_test
def test_master_mismatch_warns_but_keeps_parsed_customer(up_db):
    """파일 품번이 엠제이티 목록에 없는데 매칭된 마스터 품번이 엠제이티 전용이면 경고만, 저장은 파싱 판정."""
    up_db(existing_mjt=False)
    at = _upload([("OLD-PDVN02;OUP", "PDVN 2", 5, 100)])
    assert any("거래처 판정이 마스터 품번과 다른 라인 1개" in str(w.value) and "4PDVN-02" in str(w.value)
               for w in at.warning), [w.value for w in at.warning]
    heads, _ = _save(at)
    assert [(h["customer"], h["so_number"]) for h in heads] == [(MIJIN, PO)]


# ── 표기 통일(Data Reviewer M6): customer_for 도 vendors 정식명 ──────────────────

def test_mjt_customer_constant_shared_by_parsers():
    import utils.mijin_erp_parser as erp
    import utils.so_parser as sp
    assert erp.MJT_CUSTOMER == MJT == "(주)엠제이티" and sp.MJT_CUSTOMER is erp.MJT_CUSTOMER
    assert erp.customer_for("4PDVN-03") == MJT and erp.customer_for(None, "4PDVN-02;OUP") == MJT
    assert erp.customer_for("4HDV-VM-04") == MIJIN
    rows = [{"so_number": "E1", "so_date": None, "due_date": None, "amount": 10,
             "customer_hint": erp.customer_for("8PDVN-02")},
            {"so_number": "E1", "so_date": None, "due_date": None, "amount": 5,
             "customer_hint": erp.customer_for("4HDV-VM-04")}]
    assert sorted((h["customer"], h["so_number"]) for h in erp.group_by_order(rows)) == [
        (MJT, "E1-MJT"), (MIJIN, "E1")]
    src = open(os.path.join(ROOT, "utils", "mijin_erp_parser.py"), encoding="utf-8").read()
    assert src.count("㈜") == 1 and 'return "㈜' not in src           # 옛 표기는 상수 주석 설명 1곳만
