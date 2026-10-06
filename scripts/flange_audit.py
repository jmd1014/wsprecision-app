"""플랜지 소재 점검표 생성 — 도면 치수(DRW) × 명진 양식 × BOM 대조 → _플랜지_소재점검_YYYYMMDD.xlsx (2026-10-06, 읽기 전용)"""
import re, sys
sys.path.insert(0, r"C:\wsapp")
import openpyxl, db
from openpyxl.styles import PatternFill, Font

# 도면에서 읽은 제품 치수 (외경 × 두께) — 2026-10-06
DRW = {
 "15BLD":(95,12),"20BLD":(100,14),"25BLD":(125,14),"32BLD":(135,16),"40BLD":(140,16),"50BLD":(155,16),"65BLD":(175,18),
 "T12BLD":(95,12),"T16BLD":(100,14),"T20BLD":(125,14),"T24BLD":(135,16),"T24BLD-X156":(125,19.1),"T32BLD":(140,16),"T32BLD-X156":(150,19.1),"T48BLD":(175,18),
 "15LJF":(95,12),"20LJF":(100,14),"25LJF":(125,14),"32LJF":(135,16),"40LJF":(140,16),"50LJF":(155,16),"65LJF":(175,18),
 "T12LJF":(95,12),"T16LJF":(100,14),"T16LJF-X156":(110,12.7),"T20LJF":(125,14),"T24LJF":(135,16),"T24LJF-X156":(125,17.5),"T32LJF":(140,16),"T32LJF-X156":(150,17.5),
 "15RFF":(95,12),"15RFF-X363":(95,14),"20RFF":(100,14),"25RFF":(125,14),"32RFF":(135,16),"40RFF":(140,16),"50RFF":(155,16),"65RFF":(175,18),
 "15ABV-FL-01":(95,17),"20ABV-FL-01":(100,19),"25ABV-FL-01":(125,20),"25ABV-FL-01-X156":(108,21.9),"32ABV-FL-01":(135,17.5),"40ABV-FL-01":(140,17.5),
 "50ABV-FL-01":(155,21.6),"50ABV-FL-01-X156":(152,26.6),"T16ABV-FL-01":(100,19),"50PIF-01":(155,17.5),"T32PIF-19L-01-X171":(152.4,19.1),
 "80ABV-FL-05":(112.6,13),"100ABV-FL-06":(146,17),"150ABV-FL-05":(100,15),"200ABV-FL-06":(100,15),
}
DRW_FILE = {"T24BLD-X156":"T24BLD-X156.pdf","T32BLD-X156":"T32BLD-X156-APP.png","T16LJF-X156":"T16LJF-X156.pdf","T24LJF-X156":"T24LJF-X156.pdf",
            "T32LJF-X156":"T32LJF-X156-APP.png","15RFF-X363":"15RFF-X363-APP.png","25ABV-FL-01-X156":"25ABV-FL-01-X156.png","50ABV-FL-01-X156":"50ABV-FL-01-X156.png",
            "80ABV-FL-05":"BELLOWS FLANGE/80ABV-FL-05_01.png","100ABV-FL-06":"BELLOWS FLANGE/100ABV-FL-06_00.png",
            "150ABV-FL-05":"BELLOWS FLANGE/150ABV-FL-05_01.png","200ABV-FL-06":"BELLOWS FLANGE/200ABV-FL-06_01.jpeg"}

wb = openpyxl.load_workbook(r"G:\내 드라이브\제품관리DB\명진메탈 발주서.xlsx", data_only=True)
form = {}
for name in ["FLANGE", "B FLANGE", "R FLANGE", "P FLANGE", "J FLANGE", "미진(4)"]:
    for r in wb[name].iter_rows(values_only=True):
        vals = [v for v in r if v is not None]
        if len(vals) < 2 or not isinstance(vals[0], int):
            continue
        m = re.match(r"\s*([A-Za-z0-9\-/]+)\s*\(([^)]*)\)", str(vals[1]))
        if not m:
            continue
        base, pdim = m.group(1), m.group(2)
        grade = None
        spec = None
        for v in vals[2:]:
            sv = str(v)
            if re.fullmatch(r"3(04|16)", sv.strip()):
                grade = sv.strip()
                continue
            if re.search(r"\d+\s*\*\s*\d", sv) and spec is None:
                g = re.match(r"\s*(?:SUS|S)?\s*(304|316)\s+(.*)", sv)
                if g:
                    grade = grade or g.group(1)
                    spec = g.group(2)
                else:
                    spec = sv
        if spec is None:
            continue
        key = re.sub(r"-(APP|MP|AP|M|P)$", "", base.upper())
        form.setdefault(key, []).append({"base": base, "pdim": pdim.replace(" ", ""), "grade": grade, "spec": spec.strip().replace(" ", "")})

prods = db.fetch("products", "product_id,pn,material,customer",
                 "archived_at=is.null&or=(pn.ilike.*FL*,pn.ilike.*BLD*,pn.ilike.*LJF*,pn.ilike.*RFF*,pn.ilike.*PIF*)", limit=2000)
P = []
seen = set()
for p in prods:
    # ABV-FL-16/17 은 육각 소재 그랜드너트류 — 플랜지 아님 (2026-10-06 사용자)
    if p["pn"] in seen or not re.match(r"^4?ST?\d+(BLD|LJF|RFF|PIF|ABV)", p["pn"]) or re.search(r"ABV-FL-1[67]$", p["pn"]):
        continue
    seen.add(p["pn"])
    P.append(p)
pids = ",".join(f'"{p["product_id"]}"' for p in P)
bom = {b["product_id"]: b for b in db.fetch("bom", "product_id,material_id", f"product_id=in.({pids})&process_type=eq.MATERIAL", limit=1000)}
mids = ",".join(f'"{b["material_id"]}"' for b in bom.values() if b["material_id"])
mats = {m["material_id"]: m for m in db.fetch("materials", "material_id,raw_name", f"material_id=in.({mids})", limit=1000)}
stock = {s["material_id"]: float(s["current_stock"] or 0) for s in db.fetch("material_stock", "material_id,current_stock", f"material_id=in.({mids})", limit=1000)}
_bid2pid = {str(b["bom_id"]): b["product_id"] for b in db.fetch("bom", "bom_id,product_id", f"product_id=in.({pids})&process_type=eq.MATERIAL", limit=1000)}
chg = {_bid2pid[c["record_id"]] for c in db.fetch("master_change_log", "record_id", "table_name=eq.bom&reason=like.*2026-10-06*", limit=1000) if c["record_id"] in _bid2pid}


def base_of(pn):
    k = re.sub(r"^4", "", pn.upper())
    k = re.sub(r"^S(?=T?\d)", "", k)
    k = re.sub(r"-(APP|MP|AP|M|P)$", "", k)
    return re.sub(r"#S$", "", k)


def dims(s):
    s = str(s)
    s = s[s.find("Ø") + 1:] if "Ø" in s else s
    m = re.findall(r"(\d+(?:\.\d+)?)", s.replace("ø", ""))
    return (float(m[0]), float(m[1])) if len(m) >= 2 else None


rows = []
for p in sorted(P, key=lambda x: x["pn"]):
    b = bom.get(p["product_id"])
    mid = b["material_id"] if b else None
    mat = mats.get(mid, {})
    grade = "304" if "304" in (p["material"] or "") else "316"
    key = base_of(p["pn"])
    d = DRW.get(key)
    cands = form.get(key, [])
    f = next((c for c in cands if c["grade"] == grade), None) or (cands[0] if cands else None)
    bd = dims(mat.get("raw_name", "")) if mat else None
    note = []
    if not mid:
        st = "BOM 없음"
    elif not d:
        st = "도면 미확인"
        rn = mat.get("raw_name") or ""
        note.append("육각/특수 소재 — 도면 치수 대조 생략" if ("H" in rn or "PT" in rn) else "도면 파일 없음")
    elif not bd:
        st = "확인 필요"
        note.append("소재명에서 치수를 읽지 못함")
    else:
        od, t = d
        mo, mt = bd
        ok_od = mo >= od and mo - od <= 5.5
        ok_t = (t + 1.4) <= mt <= (t + 3.1)
        if ok_od and ok_t:
            st = "정정(10/6)" if p["product_id"] in chg else "일치"
        else:
            st = "확인 필요"
            if not ok_od:
                note.append(f"외경: 도면 Ø{od:g} vs 소재 Ø{mo:g}")
            if not ok_t:
                note.append(f"두께: 도면 {t:g} vs 소재 {mt:g} (여유 {mt - t:g})")
    if f and d:
        fd = dims(f["pdim"])
        if fd and (abs(fd[0] - d[0]) > 0.01 or abs(fd[1] - d[1]) > 0.01):
            note.append(f"명진 양식 제품치수 {f['pdim']} ≠ 도면")
        fs = dims(f["spec"])
        if fs and bd and fs != bd:
            note.append(f"명진 양식 소재 {f['spec']} ≠ BOM(대체 치수 또는 옛 값)")
    if f and f["grade"] != grade:
        note.append(f"양식은 {f['grade']} 줄만")
    rows.append({
        "품번": p["pn"], "재질": p["material"], "고객사": p.get("customer") or "",
        "도면(기준 품명)": key if d else "-",
        "도면 제품치수": f"Ø{d[0]:g}×{d[1]:g}" if d else "-",
        "표준 소재(두께+2)": f"Ø{d[0]:g}×{d[1] + 2:g}" if d else "-",
        "현재 BOM 소재": mat.get("raw_name") or "-", "자재ID": mid or "-",
        "소재 재고": stock.get(mid, 0) if mid else "",
        "명진 양식 품명": f["base"] if f else "-", "명진 양식 제품치수": f["pdim"] if f else "-",
        "명진 양식 재질": f["grade"] if f else "-", "명진 양식 소재규격": f["spec"] if f else "-",
        "판정": st, "비고": " · ".join(note)})

from collections import Counter
print(Counter(r["판정"] for r in rows))
for r in rows:
    if r["판정"] not in ("일치", "정정(10/6)") or r["비고"]:
        print(r["판정"], "|", r["품번"], "|", r["도면 제품치수"], "|", r["현재 BOM 소재"], "|", r["비고"])

out = openpyxl.Workbook()
ws = out.active
ws.title = "플랜지 소재 점검"
cols = list(rows[0].keys())
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
fill = {"정정(10/6)": "E2F0D9", "확인 필요": "FCE4D6", "BOM 없음": "FFF2CC", "도면 미확인": "EDEDED"}
for r in rows:
    ws.append([r[c] for c in cols])
    if r["판정"] in fill:
        for c in ws[ws.max_row]:
            c.fill = PatternFill("solid", fgColor=fill[r["판정"]])
for col in ws.columns:
    ws.column_dimensions[col[0].column_letter].width = max(10, min(60, max(len(str(c.value or "")) for c in col) + 2))
ws.freeze_panes = "C2"
ws.auto_filter.ref = ws.dimensions

ws2 = out.create_sheet("도면 치수 원본")
ws2.append(["기준 품명", "도면 외경", "도면 두께", "도면 파일"])
for k, (od, t) in sorted(DRW.items()):
    ws2.append([k, od, t, DRW_FILE.get(k, k + ".png/.pdf")])
ws3 = out.create_sheet("기준")
for line in [
    "소재 표준 = 도면 외경 그대로(또는 한 치수 위), 두께 = 도면 두께 + 2 (ABV-FL 는 +1.5~2.5)",
    "15A 는 Ø95, 25A 는 Ø125 가 표준 — 명진 양식의 Ø100/Ø130 은 공급사 사정에 따른 대체 치수(입고는 실규격 자재로)",
    "65A 는 Ø175×20 표준 — 양식의 20.5 / Ø180 은 대체 치수",
    "명진 옛 발주서 양식의 괄호 안은 제품 치수(소재 아님). 5/7 자동 매칭이 이를 소재로 등록한 것이 사고 원인",
    "판정: 일치 = 도면 기준 소재 적정 / 정정(10/6) = 오늘 바로잡음 / 확인 필요 = 도면과 BOM 소재가 안 맞음 / 도면 미확인 = 육각·특수 소재",
    "도면 출처: G:\\내 드라이브\\도면\\미진정밀\\FLANGE, BELLOWS FLANGE (2026-10-06 읽음)"]:
    ws3.append([line])
ws3.column_dimensions["A"].width = 110
out.save(r"G:\내 드라이브\제품관리DB\_플랜지_소재점검_20261006.xlsx")
print("saved", len(rows))
