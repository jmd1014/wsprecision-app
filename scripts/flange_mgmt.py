"""플랜지류 소재 관리 시트 생성 — G:내 드라이브제품관리DB플랜지_소재관리.xlsx (2026-10-06)
실행: cd C:wsapp && PYTHONIOENCODING=utf-8 python scripts/flange_mgmt.py
제품 사이즈가 도면과 다르면 products.product_size 를 채우고 변경 이력에 남긴다(운영 DB 쓰기)."""
import re, sys, datetime
sys.path.insert(0, r"C:\wsapp")
import openpyxl, db
from openpyxl.styles import PatternFill, Font, Alignment

# 도면 치수 재사용 (flange_audit.py 의 DRW 사전)
import os
src = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "flange_audit.py"), encoding="utf-8").read()
exec(src[src.index("DRW = {"):src.index("DRW_FILE")])

def base_of(pn):
    k = re.sub(r"^4", "", pn.upper()); k = re.sub(r"^S(?=T?\d)", "", k)
    k = re.sub(r"-(APP|MP|AP|M|P)$", "", k); return re.sub(r"#S$", "", k)

def dims(s):
    s = str(s); s = s[s.find("Ø") + 1:] if "Ø" in s else s
    m = re.findall(r"(\d+(?:\.\d+)?)", s.replace("ø", "")); return (float(m[0]), float(m[1])) if len(m) >= 2 else None

prods = db.fetch("products", "product_id,pn,material,customer,sale_price,product_size",
                 "archived_at=is.null&or=(pn.ilike.*FL*,pn.ilike.*BLD*,pn.ilike.*LJF*,pn.ilike.*RFF*,pn.ilike.*PIF*)", limit=2000)
P = []; seen = set()
for p in prods:
    # ABV-FL-16/17 은 육각 소재 그랜드너트류 — 플랜지 아님 (2026-10-06 사용자)
    if p["pn"] in seen or not re.match(r"^4?ST?\d+(BLD|LJF|RFF|PIF|ABV)", p["pn"]) or re.search(r"ABV-FL-1[67]$", p["pn"]): continue
    seen.add(p["pn"]); P.append(p)
pids = ",".join(f'"{p["product_id"]}"' for p in P)
bom = {b["product_id"]: b for b in db.fetch("bom", "product_id,material_id", f"product_id=in.({pids})&process_type=eq.MATERIAL", limit=1000)}
mids = sorted({b["material_id"] for b in bom.values() if b["material_id"]})
mq = ",".join(f'"{m}"' for m in mids)
mats = {m["material_id"]: m for m in db.fetch("materials", "material_id,raw_name,material_type", f"material_id=in.({mq})", limit=1000)}
stock = {s["material_id"]: float(s["current_stock"] or 0) for s in db.fetch("material_stock", "material_id,current_stock", f"material_id=in.({mq})", limit=1000)}

# 최근 매입 단가: ① 앱 발주 라인(자재 ID) ② 매입 원장(명진, 품목 문자열 파싱)
po_price = {}
pos = {p["po_id"]: p for p in db.fetch("purchase_orders", "po_id,po_number,po_date,status", "status=neq.CANCELLED", limit=2000)}
for it in db.fetch("purchase_order_items", "po_id,material_id,unit_price,item_name", f"material_id=in.({mq})&unit_price=gt.0", limit=5000):
    po = pos.get(it["po_id"])
    if not po: continue
    cur = po_price.get(it["material_id"])
    if not cur or po["po_date"] > cur[0]:
        po_price[it["material_id"]] = (po["po_date"], float(it["unit_price"]), f"발주 {po['po_number']}")
# 명진 원장: unit=KG, unit_price=kg 단가, qty=EA 수, weight=총 kg, amount=weight×kg단가 → EA 단가 = amount/qty
ledger = db.fetch("purchase_ledger", "trade_date,item,unit_price,qty,weight,amount", "vendor=ilike.*명진*&unit_price=gt.0&qty=gt.0&trade_date=gte.2025-01-01", limit=5000)
led = {}  # (grade, od, t) -> (date, ea_price, item, kg_price, kg_per_ea)
for r in ledger:
    m = re.search(r"STS\s*(304|316)L?\s*환봉\s*(\d+(?:\.\d+)?)\s*￠\s*\*?\s*(\d+(?:\.\d+)?)", r["item"] or "")
    if not m: continue
    key = (m.group(1), float(m.group(2)), float(m.group(3)))
    qty = float(r["qty"] or 0); amt = float(r["amount"] or 0); wt = float(r["weight"] or 0)
    if qty <= 0 or amt <= 0: continue
    cur = led.get(key)
    if not cur or r["trade_date"] > cur[0]:
        led[key] = (r["trade_date"], round(amt / qty), r["item"], float(r["unit_price"]), round(wt / qty, 2) if wt else None)

def mat_price(mid):
    """매입 원장(세금계산서) 기준 EA 단가 — (date, ea_price, 근거, kg단가, kg/EA)"""
    m = mats.get(mid, {}); rn = m.get("raw_name", "")
    g = "304" if "304" in rn else ("316" if "316" in rn else None)
    d = dims(rn)
    if not g or not d: return None
    if (g, d[0], d[1]) in led:
        v = led[(g, d[0], d[1])]; return (v[0], v[1], f"원장 {v[2]}", v[3], v[4])
    cands = [(k, v) for k, v in led.items() if k[0] == g and k[2] == d[1] and 0 < k[1] - d[0] <= 5]
    if cands:
        k, v = sorted(cands, key=lambda kv: kv[1][0], reverse=True)[0]
        return (v[0], v[1], f"원장 {v[2]} (대체 외경)", v[3], v[4])
    return None

rows = []; size_updates = []
for p in sorted(P, key=lambda x: x["pn"]):
    b = bom.get(p["product_id"]); mid = b["material_id"] if b else None
    mat = mats.get(mid, {}); key = base_of(p["pn"]); d = DRW.get(key)
    psize = f"Ø{d[0]:g}*{d[1]:g}L" if d else (p.get("product_size") or "")
    if d and (p.get("product_size") or "") != psize:
        size_updates.append((p["product_id"], p["pn"], p.get("product_size"), psize))
    sale = float(p["sale_price"]) if p.get("sale_price") not in (None, "") else None
    pr = mat_price(mid) if mid else None
    pp = po_price.get(mid) if mid else None
    mp = pr[1] if pr else (pp[1] if pp else None)
    ratio = (mp / sale) if (mp and sale) else None
    rows.append({"품번": p["pn"], "재질": p["material"], "고객사": p.get("customer") or "",
                 "제품 사이즈(도면)": psize, "소재 사이즈(BOM)": mat.get("raw_name") or "-", "자재ID": mid or "-",
                 "소재 재고": stock.get(mid, 0) if mid else None,
                 "판매 단가": sale, "소재 단가(최근 매입)": mp, "소재비 비중": ratio,
                 "kg 단가": pr[3] if pr else None, "kg/EA": pr[4] if pr else None,
                 "앱 발주 최근 단가": pp[1] if pp else None,
                 "매입 근거": (f"{pr[0]} {pr[2]}" if pr else (f"{pp[0]} {pp[2]} (원장 없음)" if pp else "")), "비고": ""})

# ── 제품 사이즈 채우기 (도면 기준) + 변경 이력 ──
today = datetime.date.today().isoformat()
for pid, pn, old, new in size_updates:
    db.update("products", f"product_id=eq.{pid}", {"product_size": new})
    db.insert("master_change_log", [{"table_name": "products", "record_id": pid, "field_name": "product_size",
                                     "old_value": old, "new_value": new, "changed_by": "Claude",
                                     "reason": f"[일괄 정정 {today}] 플랜지 제품 사이즈를 도면 치수로 채움 ({pn})"}])
print("product_size 갱신", len(size_updates))

# ── 소재별 요약 ──
summ = {}
for r in rows:
    if r["자재ID"] == "-": continue
    s = summ.setdefault(r["자재ID"], {"자재ID": r["자재ID"], "소재": r["소재 사이즈(BOM)"], "제품 수": 0, "제품": [], "소재 재고": r["소재 재고"], "최근 매입 단가": r["소재 단가(최근 매입)"], "매입 근거": r["매입 근거"]})
    s["제품 수"] += 1; s["제품"].append(r["품번"])

out = openpyxl.Workbook(); ws = out.active; ws.title = "플랜지 소재 관리"
cols = list(rows[0].keys()); ws.append(cols)
for c in ws[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7"); c.alignment = Alignment(horizontal="center")
for r in rows:
    ws.append([r[c] for c in cols])
    row = ws[ws.max_row]
    for i in (6, 7, 8, 10, 12): row[i].number_format = "#,##0"
    row[9].number_format = "0.0%"; row[11].number_format = "0.00"
    if r["소재비 비중"] is not None and r["소재비 비중"] >= 0.8:
        for c in row: c.fill = PatternFill("solid", fgColor="FCE4D6")
    elif r["판매 단가"] is None:
        row[7].fill = PatternFill("solid", fgColor="FFF2CC")
for col in ws.columns: ws.column_dimensions[col[0].column_letter].width = max(10, min(48, max(len(str(c.value or "")) for c in col) + 2))
ws.freeze_panes = "B2"; ws.auto_filter.ref = ws.dimensions
ws2 = out.create_sheet("소재별 요약"); ws2.append(["자재ID", "소재", "제품 수", "소재 재고", "최근 매입 단가", "매입 근거", "사용 제품"])
for c in ws2[1]: c.font = Font(bold=True); c.fill = PatternFill("solid", fgColor="DDEBF7")
for s in sorted(summ.values(), key=lambda x: x["소재"]):
    ws2.append([s["자재ID"], s["소재"], s["제품 수"], s["소재 재고"], s["최근 매입 단가"], s["매입 근거"], ", ".join(s["제품"])])
    ws2[ws2.max_row][4].number_format = "#,##0"
for col in ws2.columns: ws2.column_dimensions[col[0].column_letter].width = max(10, min(70, max(len(str(c.value or "")) for c in col) + 2))
ws2.freeze_panes = "A2"
ws3 = out.create_sheet("설명")
for line in ["플랜지류 소재 관리 시트 — 2026-10-06 생성 (앱 DB 기준 스냅샷; 최신 값은 앱 마스터/BOM 이 진실)",
             "제품 사이즈 = 도면(G:\\내 드라이브\\도면\\미진정밀\\FLANGE) 외경×두께. 소재 사이즈 = BOM 자재(표준 규격). 둘의 차이가 가공여유",
             "판매 단가 = 제품 마스터 sale_price (비어 있으면 노란색 — 수주 들어올 때 채우는 중)",
             "소재 단가 = (당분간) 명진 매입 원장(세금계산서, 2025~)의 최근 EA 단가 — 앱 발주·입고 이력이 쌓이면 그쪽으로 바꿀 예정(= 공급가액 ÷ 수량). 원장은 kg 단가×중량으로 청구되므로 kg 단가와 kg/EA 도 함께 표시. '(대체 외경)' 은 명진이 한 치수 큰 봉으로 공급한 값. 원장이 없으면 앱 발주 라인 단가(프리필 값이라 참고용)",
             "소재비 비중 = 소재 단가 ÷ 판매 단가. 80% 이상은 주황색 — 단가 재협상 또는 소재 규격 재검토 대상",
             "소재 재고 = 앱 재고(입고−투입). 공급사가 다른 치수로 납품하면 그 실규격 자재 LOT 로 잡히므로 여기 숫자와 다를 수 있음",
             "시트를 갱신하려면 Claude 에게 '플랜지 소재 관리 시트 갱신' 요청 — 같은 스크립트로 다시 뽑는다"]:
    ws3.append([line])
ws3.column_dimensions["A"].width = 120
out.save(r"G:\내 드라이브\제품관리DB\플랜지_소재관리.xlsx")
print("saved", len(rows), "rows;", sum(1 for r in rows if r["소재비 비중"] is not None and r["소재비 비중"] >= 0.8), "over80;",
      sum(1 for r in rows if r["판매 단가"] is None), "no sale price;", sum(1 for r in rows if r["소재 단가(최근 매입)"] is None), "no mat price")
for r in rows:
    if r["소재비 비중"] is not None and r["소재비 비중"] >= 0.8:
        print(" over80:", r["품번"], r["판매 단가"], r["소재 단가(최근 매입)"], f"{r['소재비 비중']:.0%}", r["매입 근거"])
